"""Piri dokumentumkészítő eszközei.

- Szöveges dokumentum: Markdown a forrás, letöltéskor készül belőle Word (.docx) vagy PDF.
- Táblázat: strukturált JSON-ból Excel (.xlsx), képletekkel és számformátumokkal.
- Prezentáció: strukturált JSON-ból PowerPoint (.pptx).
- Olvasás: a tárolt .xlsx / .pptx / .docx szöveggé alakítása, hogy Piri lássa a tartalmukat.

Csak sima pip-csomagok kellenek: markdown-it-py, python-docx, reportlab, openpyxl, python-pptx.
A PDF-hez a fonts/ mappában lévő DejaVu betűtípusok kellenek (az ő és ű betűk miatt).
"""
import datetime
import io
import os
import re
from html.parser import HTMLParser
from xml.sax.saxutils import escape

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PDF_MIME = "application/pdf"
MD_MIME = "text/markdown"

MAX_SHEETS, MAX_ROWS, MAX_COLS = 10, 5000, 50
MAX_SLIDES, MAX_TABLE_ROWS, MAX_TABLE_COLS = 40, 15, 8
MAX_TEXT_FOR_CHAT = 60000  # ennyi karaktert adunk át Pirinek egy-egy xlsx/pptx/docx tartalmából


class ToolError(Exception):
    """Érthető hibaüzenet, amit Piri visszakap, és ki tud javítani."""


# ---------------------------------------------------------------------------
# Eszközdefiníciók Claude számára
# ---------------------------------------------------------------------------
TOOLS = [
    {
        "name": "create_document",
        "description": (
            "Szöveges dokumentum (szerződés, indoklás, vélemény, levél, összefoglaló, jegyzet stb.) készítése "
            "Markdown formátumban. A dokumentum a projekt dokumentumtárába kerül, ahonnan a felhasználó Word, PDF "
            "vagy Markdown formátumban töltheti le. Csak akkor használd, ha a felhasználó dokumentum vagy fájl "
            "készítését kéri; egy egyszerű kérdésre a chatben válaszolj."
        ),
        "input_schema": {
            "type": "object",
            "properties": {
                "title": {"type": "string", "description": "A dokumentum címe (a fájlnév is ebből készül)."},
                "content_markdown": {
                    "type": "string",
                    "description": (
                        "A teljes dokumentum Markdownban: # címsorok, listák, **kiemelés**, táblázatok "
                        "(| oszlop | oszlop |). A címet ne ismételd meg a tartalom elején, azt az alkalmazás beteszi."
                    ),
                },
            },
            "required": ["title", "content_markdown"],
        },
    },
    {
        "name": "create_spreadsheet",
        "description": (
            "Excel táblázat (.xlsx) készítése egy vagy több munkalappal. Az 1. sor a fejléc, az adatok a 2. sortól "
            "indulnak, az A oszlop az 1. oszlop. Képletet úgy adj meg, hogy a cella értéke '='-lel kezdődő szöveg, "
            "angol függvénynevekkel és vesszős argumentum-elválasztással (pl. '=B2*C2', '=SUM(B2:B10)', "
            "'=IF(B2>0,B2*0.27,0)'); az Excel magyarul jeleníti meg. A számokat számként add meg (ne szövegként), "
            "a dátumokat 'ÉÉÉÉ-HH-NN' szövegként. Százalékot törtként (0.27) és '0.0%' formátummal. A sorok "
            "száma ne legyen több 5000-nél. A diagramot nem kell elkészíteni."
        ),
        "input_schema": {
            "type": "object",
            "properties": {
                "title": {"type": "string", "description": "A fájl neve."},
                "sheets": {
                    "type": "array",
                    "minItems": 1,
                    "items": {
                        "type": "object",
                        "properties": {
                            "name": {"type": "string", "description": "Munkalap neve (max. 31 karakter)."},
                            "columns": {
                                "type": "array",
                                "items": {
                                    "type": "object",
                                    "properties": {
                                        "header": {"type": "string"},
                                        "format": {
                                            "type": "string",
                                            "description": (
                                                "Excel számformátum az oszlop celláira, pl. '#,##0' (egész, "
                                                "ezres tagolással), '#,##0.00', '0.0%', 'yyyy-mm-dd', "
                                                "'#,##0 \"Ft\"', '@' (szöveg). Elhagyható."
                                            ),
                                        },
                                        "width": {"type": "number", "description": "Oszlopszélesség (opcionális)."},
                                    },
                                    "required": ["header"],
                                },
                            },
                            "rows": {
                                "type": "array",
                                "items": {
                                    "type": "array",
                                    "items": {"type": ["string", "number", "boolean", "null"]},
                                },
                                "description": "Az adatsorok; minden sor az oszlopok sorrendjében tartalmazza az értékeket.",
                            },
                        },
                        "required": ["name", "columns", "rows"],
                    },
                },
            },
            "required": ["title", "sheets"],
        },
    },
    {
        "name": "create_presentation",
        "description": (
            "PowerPoint prezentáció (.pptx) készítése. Dianként egy gondolat: rövid cím, legfeljebb 6-7 rövid "
            "felsorolási pont, vagy egy legfeljebb 15 soros, 8 oszlopos táblázat. Az első dia lehet címdia "
            "(csak 'title' és 'subtitle'). Az előadói jegyzet a 'notes' mezőbe kerül."
        ),
        "input_schema": {
            "type": "object",
            "properties": {
                "title": {"type": "string", "description": "A prezentáció címe (a fájlnév is ebből készül)."},
                "slides": {
                    "type": "array",
                    "minItems": 1,
                    "items": {
                        "type": "object",
                        "properties": {
                            "title": {"type": "string"},
                            "subtitle": {"type": "string", "description": "Csak címdián."},
                            "bullets": {"type": "array", "items": {"type": "string"}},
                            "table": {
                                "type": "object",
                                "properties": {
                                    "columns": {"type": "array", "items": {"type": "string"}},
                                    "rows": {
                                        "type": "array",
                                        "items": {"type": "array", "items": {"type": ["string", "number", "null"]}},
                                    },
                                },
                                "required": ["columns", "rows"],
                            },
                            "notes": {"type": "string", "description": "Előadói jegyzet."},
                        },
                        "required": ["title"],
                    },
                },
            },
            "required": ["title", "slides"],
        },
    },
]


def safe_filename(title, ext):
    stem = re.sub(r"[^\w\- ]", "", str(title or "dokumentum"), flags=re.UNICODE).strip()
    stem = re.sub(r"\s+", "_", stem)[:80] or "dokumentum"
    return f"{stem}.{ext}"


def run_tool(name, tool_input):
    """Végrehajt egy eszközhívást. Visszaad: (fájlnév, mime, bájtok). Hibánál ToolError."""
    if not isinstance(tool_input, dict):
        raise ToolError("Hibás eszközbemenet.")
    title = str(tool_input.get("title") or "").strip()
    if name == "create_document":
        body = str(tool_input.get("content_markdown") or "").strip()
        if not body:
            raise ToolError("A content_markdown üres.")
        if not re.match(r"^#\s", body) and title:
            body = f"# {title}\n\n{body}"
        return safe_filename(title, "md"), MD_MIME, (body + "\n").encode("utf-8")
    if name == "create_spreadsheet":
        return safe_filename(title, "xlsx"), XLSX_MIME, build_xlsx(tool_input)
    if name == "create_presentation":
        return safe_filename(title, "pptx"), PPTX_MIME, build_pptx(tool_input)
    raise ToolError(f"Ismeretlen eszköz: {name}")


# ---------------------------------------------------------------------------
# XLSX
# ---------------------------------------------------------------------------
_FORMAT_ALIASES = {
    "text": "@", "szöveg": "@", "szoveg": "@",
    "integer": "#,##0", "egész": "#,##0", "egesz": "#,##0",
    "number": "#,##0.00", "szám": "#,##0.00", "szam": "#,##0.00",
    "percent": "0.0%", "százalék": "0.0%", "szazalek": "0.0%",
    "date": "yyyy-mm-dd", "dátum": "yyyy-mm-dd", "datum": "yyyy-mm-dd",
    "currency": '#,##0 "Ft"', "pénz": '#,##0 "Ft"', "penz": '#,##0 "Ft"',
}
_BAD_FORMULA = re.compile(
    r"(WEBSERVICE|HYPERLINK|IMPORTDATA|IMPORTXML|IMPORTHTML|IMPORTRANGE|RTD|DDE|EXEC|CALL|REGISTER|URL)\s*\(|"
    r"https?:|file:|ftp:|\[|\\\\",
    re.IGNORECASE,
)
_ISO_DATE = re.compile(r"^\d{4}-\d{2}-\d{2}$")


def _norm_format(fmt):
    if not fmt:
        return None
    fmt = str(fmt).strip()
    return _FORMAT_ALIASES.get(fmt.lower(), fmt)


def _is_date_format(fmt):
    if not fmt:
        return False
    return bool(re.search(r"[ymd]", re.sub(r'"[^"]*"', "", fmt).lower()))


def _prepare_value(value, fmt):
    if isinstance(value, str):
        if value.startswith("="):
            if _BAD_FORMULA.search(value):
                raise ToolError(f"A képlet nem engedélyezett (külső hivatkozás vagy tiltott függvény): {value[:60]}")
            return value
        if _is_date_format(fmt) and _ISO_DATE.match(value):
            try:
                return datetime.date.fromisoformat(value)
            except ValueError:
                return value
    return value


def _sheet_name(raw, used):
    name = re.sub(r"[\[\]:*?/\\]", " ", str(raw or "")).strip().strip("'")[:31] or "Munkalap"
    base, n = name, 2
    while name.lower() in used:
        suffix = f" ({n})"
        name = base[: 31 - len(suffix)] + suffix
        n += 1
    used.add(name.lower())
    return name


def build_xlsx(spec):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    sheets = spec.get("sheets")
    if not isinstance(sheets, list) or not sheets:
        raise ToolError("Legalább egy munkalap (sheets) kell.")
    if len(sheets) > MAX_SHEETS:
        raise ToolError(f"Legfeljebb {MAX_SHEETS} munkalap lehet.")

    wb = Workbook()
    wb.remove(wb.active)
    used = set()
    side = Side(style="thin", color="BFBFBF")
    border = Border(left=side, right=side, top=side, bottom=side)

    for sh in sheets:
        cols = sh.get("columns") or []
        rows = sh.get("rows") or []
        if not cols:
            raise ToolError(f"A(z) '{sh.get('name', '')}' munkalapnak nincs oszlopa.")
        if len(cols) > MAX_COLS:
            raise ToolError(f"Legfeljebb {MAX_COLS} oszlop lehet munkalaponként.")
        if len(rows) > MAX_ROWS:
            raise ToolError(f"Legfeljebb {MAX_ROWS} sor lehet munkalaponként.")

        ws = wb.create_sheet(_sheet_name(sh.get("name"), used))
        fmts = [_norm_format(c.get("format")) for c in cols]
        lens = [len(str(c.get("header", ""))) for c in cols]

        for ci, col in enumerate(cols, 1):
            cell = ws.cell(row=1, column=ci, value=str(col.get("header", "")))
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F3864")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            cell.border = border

        for ri, row in enumerate(rows, 2):
            if not isinstance(row, list):
                raise ToolError(f"A(z) {ri - 1}. adatsor nem lista.")
            if len(row) > len(cols):
                raise ToolError(
                    f"A(z) {ri - 1}. adatsorban több érték ({len(row)}) van, mint oszlop ({len(cols)})."
                )
            for ci in range(len(cols)):
                value = _prepare_value(row[ci] if ci < len(row) else None, fmts[ci])
                cell = ws.cell(row=ri, column=ci + 1, value=value)
                if fmts[ci]:
                    cell.number_format = fmts[ci]
                cell.border = border
                is_text = isinstance(value, str) and not value.startswith("=")
                cell.alignment = Alignment(vertical="top", wrap_text=is_text)
                lens[ci] = max(lens[ci], len(str(value)) if is_text or isinstance(value, (int, float)) else 12)

        for ci, col in enumerate(cols, 1):
            width = col.get("width")
            if not isinstance(width, (int, float)) or width <= 0:
                width = min(max(lens[ci - 1] + 2, 8), 60)
            ws.column_dimensions[get_column_letter(ci)].width = width

        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(rows) + 1}"

    wb.calculation.fullCalcOnLoad = True  # a képletek megnyitáskor kiszámolódjanak
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# PPTX
# ---------------------------------------------------------------------------
def build_pptx(spec):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt

    slides = spec.get("slides")
    if not isinstance(slides, list) or not slides:
        raise ToolError("Legalább egy dia (slides) kell.")
    if len(slides) > MAX_SLIDES:
        raise ToolError(f"Legfeljebb {MAX_SLIDES} dia lehet.")

    navy, grey = RGBColor(0x1F, 0x38, 0x64), RGBColor(0x59, 0x59, 0x59)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    prs.core_properties.title = str(spec.get("title") or "")

    def set_text(shape, text, size, bold=False, color=None):
        tf = shape.text_frame
        tf.word_wrap = True
        tf.text = str(text)
        for p in tf.paragraphs:
            for r in p.runs:
                r.font.size, r.font.bold = Pt(size), bold
                if color is not None:
                    r.font.color.rgb = color

    def place(shape, left, top, width, height):
        shape.left, shape.top, shape.width, shape.height = Inches(left), Inches(top), Inches(width), Inches(height)

    def col_weights(columns, rows):
        w = []
        for i, c in enumerate(columns):
            longest = max([len(str(c))] + [len(str(r[i])) for r in rows if i < len(r) and r[i] is not None])
            w.append(max(min(longest, 40), 6))
        return w

    for idx, s in enumerate(slides):
        title = str(s.get("title") or "")
        bullets = [str(b) for b in (s.get("bullets") or [])]
        table = s.get("table")
        subtitle = s.get("subtitle")

        if idx == 0 and subtitle and not bullets and not table:
            slide = prs.slides.add_slide(prs.slide_layouts[0])
            place(slide.shapes.title, 0.8, 2.3, 11.7, 1.6)
            set_text(slide.shapes.title, title, 40, True, navy)
            place(slide.placeholders[1], 0.8, 4.1, 11.7, 1.2)
            set_text(slide.placeholders[1], subtitle, 22, False, grey)
            for p in slide.shapes.title.text_frame.paragraphs + slide.placeholders[1].text_frame.paragraphs:
                p.alignment = 1  # balra
        else:
            use_table = isinstance(table, dict)
            slide = prs.slides.add_slide(prs.slide_layouts[5 if use_table else 1])
            place(slide.shapes.title, 0.8, 0.4, 11.7, 0.95)
            set_text(slide.shapes.title, title, 30, True, navy)
            for p in slide.shapes.title.text_frame.paragraphs:
                p.alignment = 1
            bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.38), Inches(11.7), Inches(0.05))
            bar.fill.solid()
            bar.fill.fore_color.rgb = navy
            bar.line.fill.background()

            top = 1.7
            if use_table:
                cols = [str(c) for c in (table.get("columns") or [])]
                rows = table.get("rows") or []
                if not cols:
                    raise ToolError(f"A(z) {idx + 1}. dia táblázatának nincs oszlopa.")
                if len(cols) > MAX_TABLE_COLS or len(rows) > MAX_TABLE_ROWS:
                    raise ToolError(
                        f"A(z) {idx + 1}. dia táblázata túl nagy (legfeljebb {MAX_TABLE_ROWS} sor és "
                        f"{MAX_TABLE_COLS} oszlop fér el egy dián); oszd több diára."
                    )
                if bullets:
                    box = slide.shapes.add_textbox(Inches(0.8), Inches(top), Inches(11.7), Inches(0.55 * len(bullets)))
                    box.text_frame.word_wrap = True
                    for i, b in enumerate(bullets):
                        p = box.text_frame.paragraphs[0] if i == 0 else box.text_frame.add_paragraph()
                        p.text = f"• {b}"
                        for r in p.runs:
                            r.font.size = Pt(20)
                    top += 0.55 * len(bullets) + 0.2
                n_rows = len(rows) + 1
                font = 18 if n_rows <= 6 else 16 if n_rows <= 9 else 14 if n_rows <= 12 else 11
                height = min(0.5 * n_rows, 7.5 - top - 0.5)
                gs = slide.shapes.add_table(n_rows, len(cols), Inches(0.8), Inches(top), Inches(11.7), Inches(height))
                tbl = gs.table
                weights = col_weights(cols, rows)
                total = Inches(11.7)
                for i, w in enumerate(weights):
                    tbl.columns[i].width = int(total * w / sum(weights))
                for ci, c in enumerate(cols):
                    tbl.cell(0, ci).text = c
                for ri, r in enumerate(rows, 1):
                    for ci in range(len(cols)):
                        v = r[ci] if ci < len(r) and r[ci] is not None else ""
                        tbl.cell(ri, ci).text = str(v)
                for ri_, r_ in enumerate(tbl.rows):
                    for cell in r_.cells:
                        for p in cell.text_frame.paragraphs:
                            for run in p.runs:
                                run.font.size = Pt(font)
                                if ri_ == 0:
                                    run.font.bold = True
            else:
                body = slide.placeholders[1]
                place(body, 0.8, top, 11.7, 5.2)
                size = 26 if len(bullets) <= 4 else 24 if len(bullets) <= 6 else 20
                body.text_frame.word_wrap = True
                for i, b in enumerate(bullets or [""]):
                    p = body.text_frame.paragraphs[0] if i == 0 else body.text_frame.add_paragraph()
                    p.text = b
                    for r in p.runs:
                        r.font.size = Pt(size)

        if s.get("notes"):
            slide.notes_slide.notes_text_frame.text = str(s["notes"])

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Markdown -> blokkok (közös köztes forma a Word és a PDF számára)
# ---------------------------------------------------------------------------
def _md_to_html(text):
    from markdown_it import MarkdownIt

    return MarkdownIt("commonmark", {"html": False}).enable("table").render(text)


class _Blocks(HTMLParser):
    """HTML -> blokklista. Blokk: ("h", szint, runs) | ("p", runs) | ("quote", runs) |
    ("li", mélység, rendezett, sorszám, runs) | ("code", runs) | ("hr",) | ("table", sorok).
    Run: (szöveg, félkövér, dőlt, kód). Táblázatcella: (runs, fejléc, igazítás)."""

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.blocks, self.kind, self.extra, self.runs = [], None, {}, None
        self.bold = self.ital = self.code = self.quote = 0
        self.lists, self.in_pre = [], False
        self.table = self.row = self.cell = None
        self.cell_head, self.cell_align, self.href = False, None, None

    # -- segédek --
    def _start(self, kind, **extra):
        self._flush()
        self.kind, self.extra, self.runs = kind, extra, []

    def _flush(self):
        if self.kind is None:
            return
        runs = self.runs or []
        if runs and not self.in_pre:
            runs[0] = (runs[0][0].lstrip(),) + runs[0][1:]
            runs[-1] = (runs[-1][0].rstrip(),) + runs[-1][1:]
        runs = [r for r in runs if r[0]]
        k = self.kind
        if k == "h" and runs:
            self.blocks.append(("h", self.extra["level"], runs))
        elif k in ("p", "quote") and runs:
            self.blocks.append((k, runs))
        elif k == "li" and runs:
            self.blocks.append(("li", self.extra["depth"], self.extra["ordered"], self.extra["n"], runs))
        elif k == "code":
            self.blocks.append(("code", [(("".join(r[0] for r in runs)).rstrip("\n"), False, False, True)]))
        self.kind, self.runs = None, None

    def _add(self, text, bold=None, ital=None, code=None):
        run = (text, bool(self.bold) if bold is None else bold, bool(self.ital) if ital is None else ital,
               bool(self.code) if code is None else code)
        if self.cell is not None:
            self.cell.append(run)
        else:
            if self.runs is None:
                self._start("quote" if self.quote else "p")
            self.runs.append(run)

    # -- HTML események --
    def handle_starttag(self, tag, attrs):
        a = dict(attrs)
        if re.fullmatch(r"h[1-6]", tag):
            self._start("h", level=int(tag[1]))
        elif tag == "p":
            if self.cell is None and self.kind != "li":
                self._start("quote" if self.quote else "p")
            elif self.kind == "li" and self.runs:
                self._add(" ")
        elif tag in ("ul", "ol"):
            self._flush()
            self.lists.append([tag == "ol", int(a.get("start") or 1) - 1])
        elif tag == "li":
            entry = self.lists[-1] if self.lists else [False, 0]
            entry[1] += 1
            self._start("li", depth=max(len(self.lists), 1), ordered=entry[0], n=entry[1])
        elif tag in ("strong", "b"):
            self.bold += 1
        elif tag in ("em", "i"):
            self.ital += 1
        elif tag == "code" and not self.in_pre:
            self.code += 1
        elif tag == "pre":
            self._flush()
            self.in_pre, self.kind, self.runs = True, "code", []
        elif tag == "blockquote":
            self._flush()
            self.quote += 1
        elif tag == "br":
            self._add("\n")
        elif tag == "hr":
            self._flush()
            self.blocks.append(("hr",))
        elif tag == "a":
            self.href = a.get("href")
        elif tag == "table":
            self._flush()
            self.table = []
        elif tag == "tr":
            self.row = []
        elif tag in ("th", "td"):
            self.cell, self.cell_head = [], tag == "th"
            m = re.search(r"text-align:\s*(\w+)", a.get("style") or "")
            self.cell_align = m.group(1) if m else None

    def handle_endtag(self, tag):
        if re.fullmatch(r"h[1-6]", tag):
            if self.kind == "h":
                self._flush()
        elif tag == "p":
            if self.kind in ("p", "quote"):
                self._flush()
        elif tag == "li":
            if self.kind == "li":
                self._flush()
        elif tag in ("ul", "ol"):
            self._flush()
            if self.lists:
                self.lists.pop()
        elif tag in ("strong", "b"):
            self.bold = max(0, self.bold - 1)
        elif tag in ("em", "i"):
            self.ital = max(0, self.ital - 1)
        elif tag == "code" and not self.in_pre:
            self.code = max(0, self.code - 1)
        elif tag == "pre":
            self._flush()
            self.in_pre = False
        elif tag == "blockquote":
            self._flush()
            self.quote = max(0, self.quote - 1)
        elif tag == "a":
            if self.href and self.href.startswith(("http://", "https://")):
                self._add(f" ({self.href})", bold=False, ital=False, code=False)
            self.href = None
        elif tag in ("th", "td"):
            if self.row is not None and self.cell is not None:
                self.row.append((self.cell, self.cell_head, self.cell_align))
            self.cell = None
        elif tag == "tr":
            if self.table is not None and self.row is not None:
                self.table.append(self.row)
            self.row = None
        elif tag == "table":
            if self.table:
                self.blocks.append(("table", self.table))
            self.table = None

    def handle_data(self, data):
        if self.in_pre:
            self.runs.append((data, False, False, True))
            return
        text = re.sub(r"\s+", " ", data)
        if not text.strip() and (self.cell is None and not self.runs):
            return
        if text:
            self._add(text)


def _parse_markdown(text):
    p = _Blocks()
    p.feed(_md_to_html(text or ""))
    p.close()
    p._flush()
    return p.blocks


# ---------------------------------------------------------------------------
# Markdown -> DOCX
# ---------------------------------------------------------------------------
def markdown_to_docx(md_text, title=""):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    blocks = _parse_markdown(md_text)
    doc = Document()
    doc.core_properties.title = title or ""
    for section in doc.sections:
        section.left_margin = section.right_margin = Cm(2.5)
        section.top_margin = section.bottom_margin = Cm(2.2)

    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(11)
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "hu-HU")  # helyesírás-ellenőrzés magyarul
    normal.element.get_or_add_rPr().append(lang)

    def add_runs(par, runs, size=None, force_bold=False):
        for text, b, i, c in runs:
            for k, part in enumerate(text.split("\n")):
                if k:
                    par.add_run().add_break()
                if not part:
                    continue
                r = par.add_run(part)
                if b or force_bold:
                    r.bold = True
                if i:
                    r.italic = True
                if c:
                    r.font.name, r.font.size = "Consolas", Pt(9.5)
                if size and not c:
                    r.font.size = Pt(size)

    align = {"right": WD_ALIGN_PARAGRAPH.RIGHT, "center": WD_ALIGN_PARAGRAPH.CENTER}

    for blk in blocks:
        kind = blk[0]
        if kind == "h":
            add_runs(doc.add_heading(level=min(blk[1], 4)), blk[2])
        elif kind == "p":
            add_runs(doc.add_paragraph(), blk[1])
        elif kind == "quote":
            add_runs(doc.add_paragraph(style="Quote"), blk[1])
        elif kind == "li":
            _, depth, ordered, n, runs = blk
            if ordered:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.9 * depth + 0.3)
                p.paragraph_format.first_line_indent = Cm(-0.6)
                p.paragraph_format.space_after = Pt(3)
                p.add_run(f"{n}. ")
            else:
                p = doc.add_paragraph(style="List Bullet" if depth <= 1 else f"List Bullet {min(depth, 3)}")
                p.paragraph_format.space_after = Pt(3)
            add_runs(p, runs)
        elif kind == "code":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            add_runs(p, blk[1])
        elif kind == "hr":
            p = doc.add_paragraph()
            pbdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "999999")):
                bottom.set(qn(k), v)
            pbdr.append(bottom)
            p._p.get_or_add_pPr().append(pbdr)
        elif kind == "table":
            rows = blk[1]
            ncols = max(len(r) for r in rows)
            table = doc.add_table(rows=len(rows), cols=ncols)
            table.style = "Table Grid"
            for ri, row in enumerate(rows):
                for ci in range(ncols):
                    cell = table.cell(ri, ci)
                    par = cell.paragraphs[0]
                    if ci < len(row):
                        runs, head, al = row[ci]
                        add_runs(par, runs, size=10, force_bold=head)
                        if al in align:
                            par.alignment = align[al]
                        if head:
                            shd = OxmlElement("w:shd")
                            for k, v in (("w:val", "clear"), ("w:color", "auto"), ("w:fill", "D9E2F3")):
                                shd.set(qn(k), v)
                            cell._tc.get_or_add_tcPr().append(shd)
            doc.add_paragraph()

    # oldalszám a láblécben
    fp = doc.sections[0].footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    fp._p.append(fld)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Markdown -> PDF (reportlab, DejaVu betűtípussal az ő/ű miatt)
# ---------------------------------------------------------------------------
_FONTS_READY = False


def _register_fonts():
    global _FONTS_READY
    if _FONTS_READY:
        return
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    here = os.path.dirname(os.path.abspath(__file__))
    dirs = [os.path.join(here, "fonts"), "/usr/share/fonts/truetype/dejavu"]

    def find(name):
        for d in dirs:
            path = os.path.join(d, name)
            if os.path.exists(path):
                return path
        raise ToolError(f"Hiányzik a betűtípusfájl a PDF-hez: fonts/{name}")

    for reg, file in (("DejaVu", "DejaVuSans.ttf"), ("DejaVu-Bold", "DejaVuSans-Bold.ttf"),
                      ("DejaVu-Italic", "DejaVuSans-Oblique.ttf"), ("DejaVu-BoldItalic", "DejaVuSans-BoldOblique.ttf"),
                      ("DejaVuMono", "DejaVuSansMono.ttf")):
        pdfmetrics.registerFont(TTFont(reg, find(file)))
    pdfmetrics.registerFontFamily("DejaVu", normal="DejaVu", bold="DejaVu-Bold",
                                  italic="DejaVu-Italic", boldItalic="DejaVu-BoldItalic")
    _FONTS_READY = True


def markdown_to_pdf(md_text, title=""):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    _register_fonts()
    blocks = _parse_markdown(md_text)
    navy = colors.HexColor("#1F3864")

    body = ParagraphStyle("body", fontName="DejaVu", fontSize=10, leading=14.5, spaceAfter=6)
    heads = {
        1: ParagraphStyle("h1", parent=body, fontName="DejaVu-Bold", fontSize=17, leading=21, textColor=navy,
                          spaceBefore=10, spaceAfter=8, keepWithNext=1),
        2: ParagraphStyle("h2", parent=body, fontName="DejaVu-Bold", fontSize=14, leading=18, textColor=navy,
                          spaceBefore=10, spaceAfter=6, keepWithNext=1),
        3: ParagraphStyle("h3", parent=body, fontName="DejaVu-Bold", fontSize=12, leading=16, textColor=navy,
                          spaceBefore=8, spaceAfter=4, keepWithNext=1),
        4: ParagraphStyle("h4", parent=body, fontName="DejaVu-Bold", fontSize=10.5, leading=14, textColor=navy,
                          spaceBefore=6, spaceAfter=3, keepWithNext=1),
    }
    quote = ParagraphStyle("quote", parent=body, fontName="DejaVu-Italic", textColor=colors.HexColor("#555555"),
                           leftIndent=14, borderPadding=0)
    cell_styles = {a: ParagraphStyle(f"cell_{a}", parent=body, fontSize=8.5, leading=11.5, spaceAfter=0, alignment=al)
                   for a, al in (("left", 0), ("center", 1), ("right", 2))}

    def markup(runs, force_bold=False):
        out = []
        for text, b, i, c in runs:
            t = escape(text).replace("\n", "<br/>")
            if c:
                t = f'<font name="DejaVuMono" size="8.5">{t.replace("  ", "&nbsp;&nbsp;")}</font>'
            if b or force_bold:
                t = f"<b>{t}</b>"
            if i:
                t = f"<i>{t}</i>"
            out.append(t)
        return "".join(out)

    avail = A4[0] - 4 * cm
    story = []
    for blk in blocks:
        kind = blk[0]
        if kind == "h":
            story.append(Paragraph(markup(blk[2]), heads[min(blk[1], 4)]))
        elif kind == "p":
            story.append(Paragraph(markup(blk[1]), body))
        elif kind == "quote":
            story.append(Paragraph(markup(blk[1]), quote))
        elif kind == "li":
            _, depth, ordered, n, runs = blk
            style = ParagraphStyle("li", parent=body, leftIndent=18 * depth + 4, bulletIndent=18 * (depth - 1) + 4,
                                   spaceAfter=3, bulletFontName="DejaVu")
            story.append(Paragraph(markup(runs), style, bulletText=f"{n}." if ordered else "•"))
        elif kind == "code":
            story.append(Paragraph(markup(blk[1]), ParagraphStyle("code", parent=body, fontName="DejaVuMono",
                                                                  fontSize=8.5, leading=11, leftIndent=8)))
        elif kind == "hr":
            story.append(HRFlowable(width="100%", thickness=0.6, color=colors.HexColor("#999999"),
                                    spaceBefore=4, spaceAfter=8))
        elif kind == "table":
            rows = blk[1]
            ncols = max(len(r) for r in rows)
            data, weights = [], [6] * ncols
            for row in rows:
                line = []
                for ci in range(ncols):
                    if ci < len(row):
                        runs, head, al = row[ci]
                        plain = "".join(r[0] for r in runs)
                        weights[ci] = max(weights[ci], min(len(plain), 40))
                        line.append(Paragraph(markup(runs, force_bold=head), cell_styles.get(al or "left")))
                    else:
                        line.append("")
                data.append(line)
            has_head = any(c[1] for c in rows[0])
            total = sum(weights)
            tbl = Table(data, colWidths=[avail * w / total for w in weights], repeatRows=1 if has_head else 0)
            style = [("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#BFBFBF")),
                     ("VALIGN", (0, 0), (-1, -1), "TOP"),
                     ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                     ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]
            if has_head:
                style.append(("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D9E2F3")))
            tbl.setStyle(TableStyle(style))
            story.extend([tbl, Spacer(1, 8)])

    def footer(canvas, doc_):
        canvas.saveState()
        canvas.setFont("DejaVu", 8)
        canvas.setFillColor(colors.HexColor("#777777"))
        canvas.drawCentredString(A4[0] / 2, 1.1 * cm, str(doc_.page))
        canvas.restoreState()

    buf = io.BytesIO()
    pdf = SimpleDocTemplate(buf, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=2 * cm,
                            bottomMargin=2 * cm, title=title or "", author="Piri")
    pdf.build(story or [Paragraph(" ", body)], onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Office fájlok szöveggé alakítása (hogy Piri lássa a tartalmukat)
# ---------------------------------------------------------------------------
def _clip(text):
    if len(text) > MAX_TEXT_FOR_CHAT:
        return text[:MAX_TEXT_FOR_CHAT] + "\n[... a tartalom itt csonkolva ...]"
    return text


def xlsx_to_text(data):
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), data_only=False)  # a képletek szövegként maradnak meg
    out = []
    for ws in wb.worksheets:
        out.append(f"## Munkalap: {ws.title}")
        for r, row in enumerate(ws.iter_rows(values_only=True), 1):
            if r > 300:
                out.append("[... további sorok kihagyva ...]")
                break
            if any(v is not None for v in row):
                out.append(f"{r}: " + " | ".join("" if v is None else str(v) for v in row))
    return _clip("\n".join(out))


def pptx_to_text(data):
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    out = []
    for n, slide in enumerate(prs.slides, 1):
        out.append(f"## {n}. dia")
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                out.append(shape.text_frame.text.strip())
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    out.append(" | ".join(c.text for c in row.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            out.append("Jegyzet: " + slide.notes_slide.notes_text_frame.text.strip())
    return _clip("\n".join(out))


def docx_to_text(data):
    from docx import Document

    doc = Document(io.BytesIO(data))
    out = [p.text for p in doc.paragraphs if p.text.strip()]
    for t in doc.tables:
        for row in t.rows:
            out.append(" | ".join(c.text for c in row.cells))
    return _clip("\n".join(out))


def office_to_text(mime, data):
    """Visszaad szöveget, ha a mime ismert irodai formátum, különben None."""
    if mime == XLSX_MIME:
        return xlsx_to_text(data)
    if mime == PPTX_MIME:
        return pptx_to_text(data)
    if mime == DOCX_MIME:
        return docx_to_text(data)
    return None
